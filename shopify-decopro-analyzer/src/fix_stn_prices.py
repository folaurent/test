#!/usr/bin/env python3
"""
Correction des prix STN à 0 € depuis le tarif officiel + réactivation.
----------------------------------------------------------------------
Contexte : ~41 séries STN ont été importées sans prix (0 €) puis mises en
brouillon pour ne pas être vendues à 0 €. Ce script :

  1. lit le tarif STN et construit {référence -> prix TTC public} ;
  2. pour chaque produit STN, mappe chaque variante par son SKU
     (STN-<référence>) vers le prix TTC du tarif ;
  3. met à jour les prix (productVariantsBulkUpdate) ;
  4. si TOUTES les variantes d'un produit ont pu être prixées, il repasse le
     produit en ACTIVE ; sinon il le laisse en brouillon et le signale.

Aucune invention : une variante dont la référence n'est pas au tarif garde son
prix (0) et le produit reste en brouillon → listé dans le rapport.

Usage :
    DRY_RUN=1 python -m src.fix_stn_prices    # simulation + rapport
    DRY_RUN=0 python -m src.fix_stn_prices    # application réelle

.env : SHOPIFY_STORE, SHOPIFY_ADMIN_TOKEN (write_products), TARIF_XLSX
"""
from __future__ import annotations

import os
import re
import time

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

DRY_RUN = os.getenv("DRY_RUN", "1") == "1"
STORE = os.environ.get("SHOPIFY_STORE", "")
TOKEN = os.environ.get("SHOPIFY_ADMIN_TOKEN", "")
TARIF = os.getenv("TARIF_XLSX", "./data/input/LISTE_PRIX_STN_2024.xlsx")
API = f"https://{STORE}/admin/api/2024-10/graphql.json"
HEADERS = {"X-Shopify-Access-Token": TOKEN, "Content-Type": "application/json"}


def gql(query: str, variables: dict | None = None) -> dict:
    r = requests.post(API, headers=HEADERS,
                      json={"query": query, "variables": variables or {}}, timeout=30)
    r.raise_for_status()
    d = r.json()
    if d.get("errors"):
        raise RuntimeError(d["errors"])
    return d["data"]


def norm(s) -> str:
    import unicodedata
    s = "" if s is None else str(s)
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().strip().lower()


def parse_ref_prices(path: str) -> dict[str, float]:
    """{REFERENCE -> prix TTC public} depuis le tarif STN."""
    wb = load_workbook(path, read_only=True, data_only=True)
    prices: dict[str, float] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_idx = None
        for i, r in enumerate(rows):
            cells = [norm(c) for c in r]
            if any(c == "reference" or c == "ref" for c in cells) and any("ttc" in c for c in cells):
                header, header_idx = cells, i
                break
        if header_idx is None:
            continue
        def col(pred):
            for j, c in enumerate(header):
                if pred(c):
                    return j
            return None
        cref = col(lambda c: c in ("reference", "ref"))
        cttc = col(lambda c: "ttc" in c and ("public" in c or "vente" in c)) or col(lambda c: "ttc" in c)
        if cref is None or cttc is None:
            continue
        for r in rows[header_idx + 1:]:
            ref = str(r[cref]).strip() if cref < len(r) and r[cref] else ""
            raw = r[cttc] if cttc < len(r) else None
            if not ref or raw is None:
                continue
            m = re.search(r"[\d.,]+", str(raw))
            if m:
                prices[ref.upper()] = round(float(m.group().replace(",", ".")), 2)
    wb.close()
    return prices


def list_stn_products() -> list[dict]:
    """Tous les produits STN, avec variantes (id + sku + prix)."""
    out, cursor = [], None
    q = """
    query($cursor: String) {
      products(first: 40, after: $cursor, query: "vendor:'STN Ceramica'") {
        edges { cursor node {
          id title status
          variants(first: 100) { nodes { id sku price } }
        } }
        pageInfo { hasNextPage }
      }
    }"""
    while True:
        d = gql(q, {"cursor": cursor})
        edges = d["products"]["edges"]
        out += [e["node"] for e in edges]
        if not d["products"]["pageInfo"]["hasNextPage"]:
            break
        cursor = edges[-1]["cursor"]
    return out


PRICE_UPDATE = """
mutation($pid: ID!, $variants: [ProductVariantsBulkInput!]!) {
  productVariantsBulkUpdate(productId: $pid, variants: $variants) {
    userErrors { field message }
  }
}"""
ACTIVATE = """
mutation($id: ID!) {
  productUpdate(product: {id: $id, status: ACTIVE}) { userErrors { message } }
}"""


def main():
    if not STORE or not TOKEN:
        raise SystemExit("Renseigne SHOPIFY_STORE et SHOPIFY_ADMIN_TOKEN dans .env")
    print(f"=== Correction prix STN — {'SIMULATION' if DRY_RUN else 'REEL'} ===")
    ttc = parse_ref_prices(TARIF)
    print(f"Tarif : {len(ttc)} références avec prix TTC.")
    products = list_stn_products()

    repriced, reactivated, left_draft = 0, 0, []
    for p in products:
        variants = p["variants"]["nodes"]
        # ne cibler que les produits ayant au moins une variante à 0
        if not any(float(v["price"]) == 0 for v in variants):
            continue
        updates, missing = [], []
        for v in variants:
            ref = (v["sku"] or "").upper().replace("STN-", "")
            price = ttc.get(ref)
            if price is not None:
                updates.append({"id": v["id"], "price": f"{price:.2f}"})
            else:
                missing.append(v["sku"])
        if not updates:
            left_draft.append((p["title"], "aucune réf au tarif"))
            continue
        if DRY_RUN:
            print(f"[SIM] {p['title']:26} {len(updates)} variantes prixées"
                  + (f", {len(missing)} sans réf" if missing else ""))
            repriced += len(updates)
            if not missing:
                reactivated += 1
            else:
                left_draft.append((p["title"], f"{len(missing)} variantes sans prix"))
            continue
        d = gql(PRICE_UPDATE, {"pid": p["id"], "variants": updates})
        errs = d["productVariantsBulkUpdate"]["userErrors"]
        if errs:
            print(f"  ⚠️ {p['title']}: {errs}")
            continue
        repriced += len(updates)
        if not missing:                       # toutes les variantes prixées -> réactiver
            gql(ACTIVATE, {"id": p["id"]})
            reactivated += 1
            print(f"  ✅ {p['title']} — prixé + réactivé")
        else:
            left_draft.append((p["title"], f"{len(missing)} variantes sans prix"))
            print(f"  ◐ {p['title']} — partiellement prixé, laissé en brouillon")
        time.sleep(0.3)

    print("\n========== RAPPORT ==========")
    print(f"Variantes reprixées : {repriced}")
    print(f"Produits réactivés  : {reactivated}")
    print(f"Laissés en brouillon (prix incomplet / hors tarif) : {len(left_draft)}")
    for title, why in left_draft:
        print(f"  - {title} ({why})")
    if DRY_RUN:
        print("\n👉 Simulation. Mets DRY_RUN=0 pour appliquer.")


if __name__ == "__main__":
    main()
